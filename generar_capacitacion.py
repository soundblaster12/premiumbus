"""
generar_capacitacion.py — Genera el documento Word del Programa de Capacitación PremiumBus.
Ejecutar: pip install python-docx && python generar_capacitacion.py
"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
import os

OUTPUT_FILE = "Programa_Capacitacion_PremiumBus.docx"

AZUL_PREMIUM = RGBColor(0x1A, 0x3A, 0x6B)
AZUL_CLARO = RGBColor(0x2B, 0x5E, 0xA7)
GRIS_OSCURO = RGBColor(0x40, 0x40, 0x40)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)


def configurar_estilos(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = GRIS_OSCURO
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    for nivel, tam in [(1, 18), (2, 14), (3, 12)]:
        h = doc.styles[f"Heading {nivel}"]
        h.font.name = "Calibri"
        h.font.size = Pt(tam)
        h.font.color.rgb = AZUL_PREMIUM
        h.font.bold = True
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)


def agregar_portada(doc):
    for _ in range(6):
        doc.add_paragraph("")
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("🚌 PREMIUMBUS")
    r.font.size = Pt(36)
    r.font.color.rgb = AZUL_PREMIUM
    r.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("PROGRAMA DE CAPACITACIÓN A USUARIOS")
    r2.font.size = Pt(20)
    r2.font.color.rgb = AZUL_CLARO
    r2.bold = True

    doc.add_paragraph("")
    linea = doc.add_paragraph()
    linea.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rl = linea.add_run("━" * 40)
    rl.font.color.rgb = AZUL_CLARO
    rl.font.size = Pt(12)

    doc.add_paragraph("")
    datos = [
        "Sistema de Gestión de Compra de Boletos de Transporte",
        "San Luis Potosí, México",
        "",
        "Versión: 1.0",
        "Fecha: Mayo 2026",
        "Estado: Documento de Capacitación",
    ]
    for d in datos:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(d)
        rr.font.size = Pt(12)
        rr.font.color.rgb = GRIS_OSCURO
    doc.add_page_break()


def celda_color(celda, color_hex, texto, bold=True, font_color=None):
    shading = celda._element.get_or_add_tcPr()
    sh_elem = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(sh_elem)
    celda.text = ""
    p = celda.paragraphs[0]
    r = p.add_run(texto)
    r.bold = bold
    r.font.size = Pt(10)
    r.font.color.rgb = font_color or BLANCO
    r.font.name = "Calibri"


def tabla_simple(doc, encabezados, filas):
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, enc in enumerate(encabezados):
        celda_color(tabla.rows[0].cells[i], "1A3A6B", enc)
    for fila in filas:
        row = tabla.add_row()
        for i, val in enumerate(fila):
            row.cells[i].text = str(val)
            for p in row.cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.name = "Calibri"
    return tabla


def seccion_objetivo(doc):
    doc.add_heading("1. Objetivo del Programa", level=1)
    doc.add_paragraph(
        "El presente programa tiene como objetivo capacitar a los usuarios finales del sistema "
        "PremiumBus para que puedan utilizar de manera autónoma, eficiente y segura todas las "
        "funcionalidades de la aplicación móvil de gestión de compra de boletos de transporte "
        "en San Luis Potosí."
    )
    doc.add_heading("1.1 Objetivos Específicos", level=2)
    objetivos = [
        "Familiarizar al usuario con la interfaz y navegación del sistema.",
        "Enseñar el proceso completo de registro, autenticación e inicio de sesión.",
        "Capacitar en la consulta de rutas, horarios y visualización de mapas.",
        "Instruir en el proceso de compra de boletos y generación de comprobantes QR.",
        "Formar al personal administrativo en el uso del panel de administración.",
        "Garantizar que los usuarios comprendan las medidas de seguridad del sistema.",
    ]
    for o in objetivos:
        doc.add_paragraph(o, style="List Bullet")


def seccion_alcance(doc):
    doc.add_heading("2. Alcance", level=1)
    doc.add_paragraph(
        "Este programa de capacitación está dirigido a dos perfiles de usuario:"
    )
    tabla_simple(doc,
        ["Perfil", "Descripción", "Módulos"],
        [
            ["Usuario General", "Pasajeros que utilizarán la app para consultar y comprar boletos", "Módulos 1-5"],
            ["Administrador", "Personal encargado de gestionar el sistema", "Módulos 1-6"],
        ]
    )


def seccion_cronograma(doc):
    doc.add_heading("3. Cronograma de Capacitación", level=1)
    doc.add_paragraph(
        "La capacitación se divide en 3 días con sesiones teórico-prácticas:"
    )
    tabla_simple(doc,
        ["Día", "Módulo", "Tema", "Duración", "Modalidad"],
        [
            ["Día 1", "Módulo 1", "Introducción y Registro", "2 horas", "Presencial"],
            ["Día 1", "Módulo 2", "Navegación e Interfaz", "1.5 horas", "Presencial"],
            ["Día 2", "Módulo 3", "Consulta de Rutas y Mapas", "2 horas", "Presencial"],
            ["Día 2", "Módulo 4", "Compra de Boletos", "2 horas", "Presencial"],
            ["Día 3", "Módulo 5", "Perfil, Historial y QR", "1.5 horas", "Presencial"],
            ["Día 3", "Módulo 6", "Panel de Administración", "2 horas", "Presencial"],
            ["Día 3", "—", "Evaluación Final", "1 hora", "Presencial"],
        ]
    )
    doc.add_paragraph("")
    p = doc.add_paragraph()
    r = p.add_run("Duración total estimada: 12 horas (3 días)")
    r.bold = True
    r.font.color.rgb = AZUL_PREMIUM


def seccion_modulos(doc):
    doc.add_heading("4. Contenido de los Módulos", level=1)

    # Módulo 1
    doc.add_heading("Módulo 1: Introducción y Registro de Usuario", level=2)
    doc.add_heading("Objetivo:", level=3)
    doc.add_paragraph("El usuario aprenderá a instalar la aplicación y crear su cuenta.")
    doc.add_heading("Contenido teórico:", level=3)
    for item in [
        "¿Qué es PremiumBus? Presentación del sistema y sus beneficios.",
        "Requisitos del dispositivo: Android 4.0+ con conexión a internet.",
        "Instalación de la PWA desde el navegador (banner de instalación).",
        "Descripción de los datos requeridos: nombre, correo y contraseña.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Práctica guiada:", level=3)
    tabla_simple(doc,
        ["Paso", "Acción", "Resultado Esperado"],
        [
            ["1", "Abrir el navegador y acceder a la URL del sistema", "Se muestra la pantalla de Login"],
            ["2", "Tocar 'Instalar' en el banner PWA", "La app se instala en el dispositivo"],
            ["3", "Tocar '¿No tienes cuenta? Regístrate'", "Se muestra el formulario de registro"],
            ["4", "Llenar nombre, correo y contraseña", "Los campos se validan en tiempo real"],
            ["5", "Tocar 'Crear Cuenta'", "Se confirma el registro y redirige al inicio"],
        ]
    )
    doc.add_paragraph("")
    doc.add_heading("Ejercicio individual:", level=3)
    doc.add_paragraph("Cada participante creará su propia cuenta y verificará que puede iniciar sesión exitosamente.")

    # Módulo 2
    doc.add_heading("Módulo 2: Navegación e Interfaz", level=2)
    doc.add_heading("Objetivo:", level=3)
    doc.add_paragraph("El usuario conocerá la estructura de navegación y las pantallas principales.")
    doc.add_heading("Contenido teórico:", level=3)
    for item in [
        "Estructura de la aplicación: barra de navegación inferior con iconos.",
        "Pantallas principales: Inicio, Viajes, Compra, Perfil.",
        "Iconografía y significado de cada sección.",
        "Funcionamiento offline: qué funciona sin internet y qué no.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Pantallas del sistema:", level=3)
    tabla_simple(doc,
        ["Pantalla", "Icono", "Función Principal"],
        [
            ["Inicio (Home)", "🏠", "Dashboard con accesos rápidos y resumen de viajes"],
            ["Viajes", "🗺️", "Consulta de rutas disponibles con mapa interactivo"],
            ["Compra", "🎫", "Selección de viaje, asiento y proceso de compra"],
            ["Perfil", "👤", "Datos del usuario, historial y boletos con QR"],
            ["Admin", "⚙️", "Panel exclusivo para administradores"],
        ]
    )

    # Módulo 3
    doc.add_heading("Módulo 3: Consulta de Rutas y Mapas", level=2)
    doc.add_heading("Objetivo:", level=3)
    doc.add_paragraph("El usuario aprenderá a buscar rutas, consultar horarios y visualizar el recorrido en el mapa.")
    doc.add_heading("Contenido teórico:", level=3)
    for item in [
        "Listado de rutas: nombre, origen, destino, hora de salida y precio.",
        "Filtrado de rutas por nombre o destino.",
        "Mapa interactivo: marcadores de origen, destino y paradas intermedias.",
        "Información de paradas: nombre y tiempo estimado de llegada.",
        "30 rutas de transporte urbano de San Luis Potosí disponibles.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Práctica guiada:", level=3)
    tabla_simple(doc,
        ["Paso", "Acción", "Resultado Esperado"],
        [
            ["1", "Navegar a la sección 'Viajes'", "Se muestra el listado de rutas"],
            ["2", "Buscar 'Tangamanga' en el filtro", "Se filtran rutas que pasan por Tangamanga"],
            ["3", "Tocar una ruta para ver detalles", "Se muestra el mapa con el recorrido"],
            ["4", "Interactuar con el mapa (zoom, paneo)", "El mapa responde a gestos táctiles"],
            ["5", "Revisar paradas intermedias", "Se ven los marcadores y tiempos"],
        ]
    )

    # Módulo 4
    doc.add_heading("Módulo 4: Compra de Boletos", level=2)
    doc.add_heading("Objetivo:", level=3)
    doc.add_paragraph("El usuario aprenderá el proceso completo de compra de un boleto de transporte.")
    doc.add_heading("Contenido teórico:", level=3)
    for item in [
        "Selección de ruta y fecha de viaje.",
        "Mapa de asientos: asientos disponibles (azul) vs ocupados (gris).",
        "Confirmación de compra y generación de comprobante.",
        "Restricción: no se pueden comprar boletos si ya se tiene un viaje activo.",
        "Compra offline: el sistema guarda la compra y la sincroniza al reconectar.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Práctica guiada:", level=3)
    tabla_simple(doc,
        ["Paso", "Acción", "Resultado Esperado"],
        [
            ["1", "Navegar a 'Compra' desde el menú", "Se muestra el formulario de compra"],
            ["2", "Seleccionar una ruta del listado", "Se cargan los datos de la ruta"],
            ["3", "Seleccionar un asiento disponible (azul)", "El asiento se marca como seleccionado"],
            ["4", "Revisar el resumen de compra", "Se muestra ruta, asiento y precio"],
            ["5", "Confirmar la compra", "Se genera el boleto y aparece el comprobante"],
        ]
    )

    # Módulo 5
    doc.add_heading("Módulo 5: Perfil, Historial y Código QR", level=2)
    doc.add_heading("Objetivo:", level=3)
    doc.add_paragraph("El usuario aprenderá a gestionar su perfil, consultar el historial y usar sus boletos QR.")
    doc.add_heading("Contenido teórico:", level=3)
    for item in [
        "Edición del nombre de usuario desde el perfil.",
        "Sección 'Viaje Activo': visualización del boleto actual con seguimiento GPS.",
        "Sección 'Historial': archivo de viajes completados.",
        "Código QR del boleto: contiene datos del viaje para verificación.",
        "Función 'En Vivo': seguimiento GPS simulado del recorrido activo.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Práctica guiada:", level=3)
    tabla_simple(doc,
        ["Paso", "Acción", "Resultado Esperado"],
        [
            ["1", "Navegar a 'Perfil'", "Se muestra el perfil con datos del usuario"],
            ["2", "Editar el nombre y guardar", "El nombre se actualiza correctamente"],
            ["3", "Verificar el viaje activo", "Se muestra el boleto comprado con QR"],
            ["4", "Tocar 'Ver QR' en el boleto", "Se despliega el código QR del viaje"],
            ["5", "Tocar 'En Vivo'", "Se abre el mapa con seguimiento GPS simulado"],
        ]
    )

    # Módulo 6
    doc.add_heading("Módulo 6: Panel de Administración (Solo Administradores)", level=2)
    doc.add_heading("Objetivo:", level=3)
    doc.add_paragraph("El administrador aprenderá a gestionar usuarios y supervisar las operaciones del sistema.")
    doc.add_heading("Contenido teórico:", level=3)
    for item in [
        "Acceso al panel: credenciales de administrador (admin@premiumbus.com).",
        "Listado de usuarios registrados y sus roles.",
        "Supervisión de compras realizadas.",
        "Gestión de rutas y disponibilidad.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Credenciales de administrador por defecto:", level=3)
    tabla_simple(doc,
        ["Campo", "Valor"],
        [
            ["Correo", "admin@premiumbus.com"],
            ["Contraseña", "admin123"],
        ]
    )
    doc.add_paragraph("")
    p = doc.add_paragraph()
    r = p.add_run("⚠️ IMPORTANTE: Cambiar la contraseña del administrador después de la primera sesión.")
    r.bold = True
    r.font.color.rgb = RGBColor(0xEF, 0x44, 0x44)


def seccion_evaluacion(doc):
    doc.add_heading("5. Evaluación de la Capacitación", level=1)
    doc.add_paragraph(
        "Al finalizar el programa, se aplicará una evaluación práctica para verificar "
        "que los usuarios dominan las funcionalidades del sistema."
    )
    doc.add_heading("5.1 Evaluación Práctica", level=2)
    doc.add_paragraph("Cada participante deberá completar las siguientes tareas sin asistencia:")
    tabla_simple(doc,
        ["#", "Tarea", "Criterio de Éxito", "Puntos"],
        [
            ["1", "Registrarse en el sistema", "Cuenta creada exitosamente", "15"],
            ["2", "Iniciar sesión con sus credenciales", "Acceso concedido al Home", "10"],
            ["3", "Buscar una ruta específica", "La ruta se muestra en el listado", "15"],
            ["4", "Visualizar el recorrido en el mapa", "Mapa con marcadores visible", "10"],
            ["5", "Comprar un boleto seleccionando asiento", "Comprobante generado", "20"],
            ["6", "Consultar el boleto QR en su perfil", "QR visible y legible", "15"],
            ["7", "Navegar entre todas las secciones", "Todas las pantallas accesibles", "15"],
        ]
    )
    doc.add_paragraph("")
    doc.add_heading("5.2 Escala de Evaluación", level=2)
    tabla_simple(doc,
        ["Rango", "Calificación", "Resultado"],
        [
            ["90-100", "Excelente", "Aprobado — Usuario autónomo"],
            ["70-89", "Bueno", "Aprobado — Requiere práctica adicional"],
            ["50-69", "Regular", "Requiere refuerzo en módulos específicos"],
            ["0-49", "Insuficiente", "Requiere repetir la capacitación"],
        ]
    )


def seccion_requisitos(doc):
    doc.add_heading("6. Requisitos Técnicos", level=1)
    doc.add_heading("6.1 Para los Participantes", level=2)
    for item in [
        "Dispositivo móvil con Android 4.0 o superior.",
        "Conexión a internet WiFi (se proporcionará durante la capacitación).",
        "Navegador web actualizado (Chrome recomendado).",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("6.2 Para el Instructor", level=2)
    for item in [
        "Computadora con proyector para demostración.",
        "Servidor con XAMPP/WAMP configurado y base de datos cargada.",
        "Acceso a la red local donde está desplegado el sistema.",
        "Copias impresas de la guía rápida de referencia.",
    ]:
        doc.add_paragraph(item, style="List Bullet")


def seccion_soporte(doc):
    doc.add_heading("7. Soporte Post-Capacitación", level=1)
    doc.add_paragraph(
        "Después de la capacitación, los usuarios contarán con los siguientes recursos de apoyo:"
    )
    tabla_simple(doc,
        ["Recurso", "Descripción", "Disponibilidad"],
        [
            ["Guía Rápida", "Documento PDF con pasos resumidos", "USB entregado"],
            ["Soporte Técnico", "Contacto con el equipo de desarrollo", "Lunes a Viernes 9-18h"],
            ["FAQ", "Preguntas frecuentes integradas en la app", "24/7 dentro de la app"],
            ["Reinstalación", "Ejecutable en USB para reinstalar el sistema", "USB entregado"],
        ]
    )


def seccion_guia_rapida(doc):
    doc.add_heading("8. Guía Rápida de Referencia", level=1)
    doc.add_paragraph("Resumen de las acciones más frecuentes para consulta rápida:")

    doc.add_heading("Registro:", level=3)
    doc.add_paragraph("Abrir app → Tocar 'Regístrate' → Llenar datos → 'Crear Cuenta'")

    doc.add_heading("Iniciar Sesión:", level=3)
    doc.add_paragraph("Abrir app → Ingresar correo y contraseña → 'Iniciar Sesión'")

    doc.add_heading("Consultar Rutas:", level=3)
    doc.add_paragraph("Menú inferior → 'Viajes' → Buscar/seleccionar ruta → Ver mapa")

    doc.add_heading("Comprar Boleto:", level=3)
    doc.add_paragraph("Menú inferior → 'Compra' → Seleccionar ruta → Elegir asiento → Confirmar")

    doc.add_heading("Ver Boleto QR:", level=3)
    doc.add_paragraph("Menú inferior → 'Perfil' → Viaje Activo → 'Ver QR'")

    doc.add_heading("Seguimiento En Vivo:", level=3)
    doc.add_paragraph("Perfil → Viaje Activo → 'En Vivo' → Ver mapa GPS")


def seccion_firmas(doc):
    doc.add_page_break()
    doc.add_heading("9. Firmas de Conformidad", level=1)
    doc.add_paragraph(
        "Con la firma de este documento, los abajo firmantes confirman haber recibido "
        "la capacitación completa del sistema PremiumBus."
    )
    doc.add_paragraph("")
    doc.add_paragraph("")

    tabla = doc.add_table(rows=4, cols=3)
    tabla.style = "Table Grid"
    encabezados = ["Nombre del Participante", "Firma", "Fecha"]
    for i, enc in enumerate(encabezados):
        celda_color(tabla.rows[0].cells[i], "1A3A6B", enc)
    for row_idx in range(1, 4):
        for col_idx in range(3):
            tabla.rows[row_idx].cells[col_idx].text = ""
            # Set row height for signature space
            tabla.rows[row_idx].height = Cm(2)

    doc.add_paragraph("")
    doc.add_paragraph("")

    # Firma del instructor
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("_" * 40)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Nombre y Firma del Instructor")
    r.bold = True
    r.font.color.rgb = AZUL_PREMIUM

    doc.add_paragraph("")
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.add_run("_" * 40)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p4.add_run("Fecha de Capacitación")
    r2.bold = True
    r2.font.color.rgb = AZUL_PREMIUM


def main():
    doc = Document()

    # Configurar márgenes
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    configurar_estilos(doc)
    agregar_portada(doc)
    seccion_objetivo(doc)
    seccion_alcance(doc)
    seccion_cronograma(doc)
    seccion_modulos(doc)
    seccion_evaluacion(doc)
    seccion_requisitos(doc)
    seccion_soporte(doc)
    seccion_guia_rapida(doc)
    seccion_firmas(doc)

    output_path = os.path.join(os.path.dirname(__file__), OUTPUT_FILE)
    doc.save(output_path)
    print(f"✅ Documento generado: {output_path}")


if __name__ == "__main__":
    main()
