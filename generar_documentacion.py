import sys
import os
import subprocess
import time

def instalar_dependencia():
    print("Verificando dependencias...")
    try:
        import docx
    except ImportError:
        print("El paquete 'python-docx' no está instalado. Instalando ahora...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        print("Dependencia instalada correctamente. Reiniciando script...\n")
        os.execv(sys.executable, ['python'] + sys.argv)

instalar_dependencia()

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, ns

# =====================================================================
# CONFIGURACIÓN DEL DOCUMENTO (Puedes modificar estos valores)
# =====================================================================

# 1. SALTO DE HOJA POR CAPÍTULO:
# Cambia esto a False si quieres QUITAR el espacio en blanco y hacer que
# los capítulos continúen en la misma hoja sin saltar a una nueva.
SALTAR_HOJA_POR_CAPITULO = True

# 2. VOLUMEN DE TEXTO (Para alcanzar las 40 hojas):
# Valor de 2 = Aprox. 40 hojas. 
# Si quieres menos hojas, ponlo en 1. Si quieres más, ponlo en 3.
MULTIPLICADOR_DE_TEXTO = 2

# =====================================================================

def crear_elemento(name):
    return OxmlElement(name)

def crear_atributo(element, name, value):
    element.set(ns.qn(name), value)

def agregar_numero_pagina(run):
    fldChar1 = crear_elemento('w:fldChar')
    crear_atributo(fldChar1, 'w:fldCharType', 'begin')
    instrText = crear_elemento('w:instrText')
    crear_atributo(instrText, 'xml:space', 'preserve')
    instrText.text = "PAGE"
    fldChar2 = crear_elemento('w:fldChar')
    crear_atributo(fldChar2, 'w:fldCharType', 'separate')
    fldChar3 = crear_elemento('w:fldChar')
    crear_atributo(fldChar3, 'w:fldCharType', 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def agregar_indice_automatico_general(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar1 = crear_elemento('w:fldChar')
    crear_atributo(fldChar1, 'w:fldCharType', 'begin')
    instrText = crear_elemento('w:instrText')
    crear_atributo(instrText, 'xml:space', 'preserve')
    instrText.text = 'TOC \\o "1-2" \\h \\z \\u'
    fldChar2 = crear_elemento('w:fldChar')
    crear_atributo(fldChar2, 'w:fldCharType', 'separate')
    fldChar3 = crear_elemento('w:fldChar')
    crear_atributo(fldChar3, 'w:fldCharType', 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def agregar_marcador(paragraph, nombre_marcador):
    p = paragraph._p
    start = crear_elemento('w:bookmarkStart')
    crear_atributo(start, 'w:id', str(abs(hash(nombre_marcador)) % 10000))
    crear_atributo(start, 'w:name', nombre_marcador)
    p.insert(0, start)
    end = crear_elemento('w:bookmarkEnd')
    crear_atributo(end, 'w:id', str(abs(hash(nombre_marcador)) % 10000))
    p.append(end)

def agregar_enlace_interno(paragraph, texto, nombre_marcador):
    hyperlink = crear_elemento('w:hyperlink')
    crear_atributo(hyperlink, 'w:anchor', nombre_marcador)
    new_run = crear_elemento('w:r')
    rPr = crear_elemento('w:rPr')
    color = crear_elemento('w:color')
    crear_atributo(color, 'w:val', '0563C1') # Azul de link
    u = crear_elemento('w:u')
    crear_atributo(u, 'w:val', 'single')
    rPr.append(color)
    rPr.append(u)
    t = crear_elemento('w:t')
    t.text = texto
    new_run.append(rPr)
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def agregar_tabla_dinamica(doc, numero, titulo, cabeceras, filas_datos):
    p = doc.add_paragraph()
    # Añadimos el marcador oculto para que el enlace del índice nos traiga aquí
    agregar_marcador(p, f"tabla_{numero}")
    p.add_run(f"Tabla {numero}: {titulo}").bold = True
    
    table = doc.add_table(rows=1, cols=len(cabeceras))
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    for i, cabecera in enumerate(cabeceras):
        hdr_cells[i].text = str(cabecera)
        for p_cell in hdr_cells[i].paragraphs:
            for run in p_cell.runs:
                run.bold = True
                
    for row_data in filas_datos:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
            
    doc.add_paragraph("") # Espacio orgánico debajo de la tabla

def generar_documento():
    doc = Document()
    
    # ESTILOS GLOBALES
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_normal.paragraph_format.line_spacing = 1.5

    style_h1 = doc.styles['Heading 1']
    style_h1.font.name = 'Arial'
    style_h1.font.size = Pt(16)
    style_h1.font.bold = True
    style_h1.font.color.rgb = None
    
    # APLICAR CONFIGURACIÓN DE SALTO DE HOJA DEL USUARIO
    if SALTAR_HOJA_POR_CAPITULO:
        style_h1.paragraph_format.page_break_before = True
    else:
        style_h1.paragraph_format.page_break_before = False

    style_h2 = doc.styles['Heading 2']
    style_h2.font.name = 'Arial'
    style_h2.font.size = Pt(14)
    style_h2.font.bold = True
    style_h2.font.color.rgb = None

    # PIE DE PÁGINA (Solo número)
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        agregar_numero_pagina(p.add_run())

    # PRIMERAS HOJAS
    doc.add_page_break() # Hoja 1: Blanco

    # Hoja 2: Índice General
    doc.add_paragraph("Índice de Contenidos", style='Heading 1')
    doc.add_paragraph("NOTA: Para ver el índice general actualizado, haz clic derecho aquí y selecciona 'Actualizar campos'.")
    agregar_indice_automatico_general(doc)
    doc.add_page_break()

    # Hoja 3: Índice de Tablas (Clicable 100% interactivo mediante Hyperlinks internos)
    doc.add_paragraph("Índice de Tablas", style='Heading 1')
    doc.add_paragraph("Dale clic a cualquier enlace azul para navegar directamente a la tabla correspondiente:", style='Normal')
    
    tablas_lista = [
        (1, "Obligaciones y Penalizaciones del Proyecto"),
        (2, "Roles, Permisos y Actores del Sistema"),
        (3, "Objetivos y Métricas de Éxito"),
        (4, "Estado de Interfaces de Usuario (UI/UX)"),
        (5, "Tecnologías y Estándares de Codificación"),
        (6, "Matrices de Cobertura de Pruebas Unitarias"),
        (7, "Resultados de Integración de Endpoints"),
        (8, "Especificaciones de Hardware del Servidor"),
        (9, "Cronograma de Ejecución (Ene-May 2026)")
    ]
    
    for num, titulo in tablas_lista:
        p = doc.add_paragraph()
        agregar_enlace_interno(p, f"Tabla {num}. {titulo}", f"tabla_{num}")
        
    doc.add_page_break()

    # CONTENIDO ORGÁNICO
    pages_data = [
        {"cap": "Capítulo 1. Contrato de Software", "subs": ["1.1 Declaraciones", "1.2 Objeto"], "table": None},
        {"cap": None, "subs": ["1.3 Obligaciones Dev"], "table": None},
        {"cap": None, "subs": ["1.4 Obligaciones Cliente"], 
         "table": {"num": 1, "title": "Obligaciones y Penalizaciones", "headers": ["Obligación", "Descripción", "Impacto"], "rows": [["Acceso Infra", "Proveer servidores", "Retraso"], ["Validación UAT", "Aprobar entregables", "Pausa sprints"]]}},
        {"cap": None, "subs": ["1.5 Propiedad Intelectual", "1.6 Pagos"], "table": None},
        {"cap": "Capítulo 2. Requerimientos", "subs": ["2.1 Entrevistas"], "table": None},
        {"cap": None, "subs": ["2.2 Historias de Usuario"], "table": None},
        {"cap": None, "subs": ["2.3 Actores"], 
         "table": {"num": 2, "title": "Roles, Permisos y Actores", "headers": ["Actor", "Descripción", "Acceso"], "rows": [["Administrador", "Gestión total", "Total"], ["Usuario", "Consulta rutas", "Lectura"], ["Soporte", "Atención", "Avanzada"]]}},
        {"cap": None, "subs": ["2.4 Req. Funcionales"], "table": None},
        {"cap": None, "subs": ["2.5 Req. No Funcionales"], "table": None},
        {"cap": "Capítulo 3. Análisis de Necesidades", "subs": ["3.1 Problema"], "table": None},
        {"cap": None, "subs": ["3.2 Objetivos"], 
         "table": {"num": 3, "title": "Objetivos y Métricas", "headers": ["Objetivo", "Métrica", "Prioridad"], "rows": [["Disponibilidad", "Uptime 99.9%", "Alta"], ["Respuesta", "< 200ms", "Media"], ["Seguridad", "AES-256", "Alta"]]}},
        {"cap": None, "subs": ["3.3 Alcance"], "table": None},
        {"cap": None, "subs": ["3.4 Viabilidad"], "table": None},
        {"cap": None, "subs": ["3.5 Impacto"], "table": None},
        {"cap": "Capítulo 4. Diseño del Sistema", "subs": ["4.1 Arquitectura"], "table": None},
        {"cap": None, "subs": ["4.2 UI/UX"], 
         "table": {"num": 4, "title": "Estado de Interfaces (UI/UX)", "headers": ["Interfaz", "Componente", "Estado"], "rows": [["Dashboard", "Gráficas", "Completado"], ["Buscador", "Mapa", "Pruebas"], ["Perfil", "Historial", "Completado"]]}},
        {"cap": None, "subs": ["4.3 Base de Datos"], "table": None},
        {"cap": None, "subs": ["4.4 Diagramas"], "table": None},
        {"cap": None, "subs": ["4.5 APIs"], "table": None},
        {"cap": "Capítulo 5. Desarrollo y Código", "subs": ["5.1 Entorno"], "table": None},
        {"cap": None, "subs": ["5.2 Estándares"], 
         "table": {"num": 5, "title": "Tecnologías y Estándares", "headers": ["Tecnología", "Uso", "Linter"], "rows": [["Node.js", "API", "ESLint"], ["React", "UI", "Prettier"], ["MySQL", "BD", "3NF"]]}},
        {"cap": None, "subs": ["5.3 Backend"], "table": None},
        {"cap": None, "subs": ["5.4 Frontend", "5.5 Servicios"], "table": None},
        {"cap": "Capítulo 6. Elaboración Pruebas", "subs": ["6.1 Estrategia"], "table": None},
        {"cap": None, "subs": ["6.2 Matrices"], 
         "table": {"num": 6, "title": "Matrices de Cobertura", "headers": ["Módulo", "Casos", "Cobertura"], "rows": [["Auth", "45", "98%"], ["Pagos", "72", "95%"], ["GPS", "30", "100%"]]}},
        {"cap": None, "subs": ["6.3 Pruebas Unitarias"], "table": None},
        {"cap": None, "subs": ["6.4 Pruebas Carga"], "table": None},
        {"cap": None, "subs": ["6.5 Entornos QA"], "table": None},
        {"cap": "Capítulo 7. Ejecución Pruebas", "subs": ["7.1 Ejecución Unitaria"], "table": None},
        {"cap": None, "subs": ["7.2 Integración"], 
         "table": {"num": 7, "title": "Resultados Integración", "headers": ["Endpoint", "Prueba", "Resultado"], "rows": [["/login", "Integración", "Aprobado"], ["/search", "Stress", "Aprobado"], ["/payments", "Seguridad", "Aprobado"]]}},
        {"cap": None, "subs": ["7.3 Benchmarking"], "table": None},
        {"cap": None, "subs": ["7.4 UAT", "7.5 Incidencias"], "table": None},
        {"cap": "Capítulo 8. Implementación", "subs": ["8.1 Despliegue"], "table": None},
        {"cap": None, "subs": ["8.2 Entorno Operativo"], 
         "table": {"num": 8, "title": "Especificaciones Hardware", "headers": ["Entorno", "Specs", "OS"], "rows": [["Prod", "8 vCPU, 16GB", "Ubuntu"], ["BD", "16 vCPU, 32GB", "Ubuntu"], ["QA", "4 vCPU, 8GB", "Ubuntu"]]}},
        {"cap": None, "subs": ["8.3 Migración"], "table": None},
        {"cap": None, "subs": ["8.4 Capacitación"], "table": None},
        {"cap": None, "subs": ["8.5 Manuales"], "table": None},
        {"cap": "Capítulo 9. Bitácora", "subs": ["9.1 Registro", "9.2 Hitos"], "table": None},
        {"cap": None, "subs": ["9.3 Control Cambios", "9.4 Problemas"], "table": None},
        {"cap": None, "subs": ["9.5 Conclusiones"], "table": None},
        {"cap": None, "subs": ["9.6 Cronograma Final"], 
         "table": {"num": 9, "title": "Cronograma de Ejecución", "headers": ["Fase", "Inicio", "Fin", "Estado"], "rows": [["Requerimientos", "26 Ene 2026", "10 Feb 2026", "Completado"], ["Análisis", "11 Feb 2026", "05 Mar 2026", "Completado"], ["Desarrollo", "06 Mar 2026", "15 Abr 2026", "Completado"], ["Pruebas", "16 Abr 2026", "05 May 2026", "Completado"], ["Implementación", "06 May 2026", "20 May 2026", "Completado"], ["Cierre", "21 May 2026", "27 May 2026", "En Proceso"]]}}
    ]

    texto_relleno = "La implementación de estos procesos se realiza respetando íntegramente las políticas de desarrollo acordadas, garantizando una correcta trazabilidad a lo largo del ciclo de vida del software. Se aplicaron revisiones exhaustivas y pruebas iterativas para validar el rendimiento, escalabilidad y la correcta adaptación de los módulos a los requisitos del cliente final."
    texto_relleno_2 = "Además, la integración continua de los entregables ha permitido mitigar riesgos arquitectónicos, fomentando un entorno de trabajo colaborativo que optimiza los recursos tecnológicos. Cada componente documentado aquí refleja la robustez del sistema frente a cargas críticas y variaciones de entorno."

    for page in pages_data:
        if page['cap']:
            doc.add_paragraph(page['cap'], style='Heading 1')
        
        for sub in page['subs']:
            doc.add_paragraph(sub, style='Heading 2')
            
            # Ajustado para comprimir el texto y acercarse a las 40 hojas (gobernado por el multiplicador)
            for _ in range(MULTIPLICADOR_DE_TEXTO):
                doc.add_paragraph(texto_relleno)
                doc.add_paragraph(texto_relleno_2)
            
        if page['table']:
            t = page['table']
            agregar_tabla_dinamica(doc, t['num'], t['title'], t['headers'], t['rows'])

    archivo_salida = "Documentacion_Sistema_PremiumBus.docx"
    
    try:
        doc.save(archivo_salida)
        print("="*60)
        print(f"✅ ¡ÉXITO! Se generó el documento '{archivo_salida}'.")
        print("="*60)
    except PermissionError:
        print("\\n" + "="*60)
        print("⚠️ ERROR: El archivo de Word está ABIERTO.")
        print("Guardando con un nombre alternativo automáticamente...")
        nuevo_archivo = f"Documentacion_Sistema_PremiumBus_v_{int(time.time())}.docx"
        doc.save(nuevo_archivo)
        print(f"✅ ¡SOLUCIONADO! Se generó exitosamente como '{nuevo_archivo}'")
        print("="*60)

if __name__ == "__main__":
    generar_documento()
