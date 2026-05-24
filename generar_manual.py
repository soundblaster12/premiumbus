"""
Genera el Manual de Usuario PremiumBus v5.0 en formato Word (.docx)
Requiere: pip install python-docx
"""
import os, sys
try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
except ImportError:
    print("Instalando python-docx...")
    os.system(f'"{sys.executable}" -m pip install python-docx')
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(6)

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Arial'
    hs.font.color.rgb = RGBColor(26, 58, 107)

section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.54)
section.right_margin = Cm(2.54)

# Numeracion de pagina (pie derecho)
footer = section.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = fp.add_run()
fld = run._r
fldChar1 = fld.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'begin'})
fld.append(fldChar1)
instrText = fld.makeelement(qn('w:instrText'), {})
instrText.text = ' PAGE '
fld.append(instrText)
fldChar2 = fld.makeelement(qn('w:fldChar'), {qn('w:fldCharType'): 'end'})
fld.append(fldChar2)

def add_blank():
    doc.add_paragraph()
    doc.add_page_break()

def title(text, level=1):
    doc.add_heading(text, level=level)

def para(text, bold=False, italic=False, size=12):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def numbered(text):
    doc.add_paragraph(text, style='List Number')

def spacer():
    doc.add_paragraph()

# ══════════════════════════════════════════════
# PAGINA 1: EN BLANCO
# ══════════════════════════════════════════════
add_blank()

# ══════════════════════════════════════════════
# PAGINA 2: PORTADA
# ══════════════════════════════════════════════
spacer(); spacer(); spacer()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PremiumBus')
r.font.size = Pt(36); r.bold = True; r.font.color.rgb = RGBColor(26, 58, 107)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Manual de Usuario v5.0')
r2.font.size = Pt(20); r2.font.color.rgb = RGBColor(100, 100, 100)
spacer()
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run('Sistema de Transporte Publico\nSan Luis Potosi, Mexico\n\nMayo 2026')
r3.font.size = Pt(14); r3.font.color.rgb = RGBColor(120, 120, 120)
doc.add_page_break()

# ══════════════════════════════════════════════
# PAGINA 3: INDICE
# ══════════════════════════════════════════════
title('Indice de Contenido')
items = [
    '1. Introduccion al Sistema ........................... 4',
    '2. Arquitectura y Documentacion Tecnica .............. 5',
    '3. Instalacion en PC (Windows) ....................... 6',
    '4. Instalacion en Android ............................ 7',
    '5. Instalacion en iOS (iPhone/iPad) .................. 8',
    '6. Guia de Uso - Usuario ............................. 9',
    '7. Guia de Uso - Administrador ....................... 10',
    '8. Funcionalidades v5.0 .............................. 11',
    '9. Recuperacion de Base de Datos ..................... 12',
    '10. Preguntas Frecuentes ............................ 13',
]
for item in items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(4)
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 1: INTRODUCCION
# ══════════════════════════════════════════════
title('1. Introduccion al Sistema')
para('PremiumBus es una plataforma digital de transporte publico disenada para la ciudad de San Luis Potosi, Mexico. Permite a los usuarios consultar rutas de autobus, visualizar paradas en un mapa interactivo, comprar boletos digitales con QR y rastrear unidades en tiempo real.')
spacer()
title('1.1 Objetivo del Sistema', level=2)
para('Proporcionar una herramienta moderna, accesible y eficiente que conecte a los pasajeros con el sistema de transporte urbano, mejorando la experiencia de viaje mediante tecnologia movil y web.')
title('1.2 Alcance', level=2)
bullet('Consulta de 30 rutas de transporte en San Luis Potosi')
bullet('Compra de boletos digitales con asiento asignado')
bullet('Generacion de codigos QR para validacion de boletos')
bullet('Rastreo GPS en tiempo real de las unidades')
bullet('Panel de administracion con reportes y analiticas')
bullet('Exportacion de datos a Excel para auditorias')
bullet('Gestion completa de conductores y usuarios')
title('1.3 Usuarios del Sistema', level=2)
bullet('Usuario Final: Pasajeros que compran boletos y consultan rutas')
bullet('Administrador: Personal autorizado que gestiona rutas, conductores, usuarios y reportes')
bullet('Administrador Principal (ID 1): Privilegios extendidos para eliminar usuarios y administradores')
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 2: DOCUMENTACION TECNICA
# ══════════════════════════════════════════════
title('2. Arquitectura y Documentacion Tecnica')
para('PremiumBus es una Progressive Web App (PWA) construida con tecnologias web estandar, compatible con navegadores modernos y capaz de funcionar como aplicacion nativa en dispositivos moviles.')
title('2.1 Tecnologias Utilizadas', level=2)
table = doc.add_table(rows=8, cols=2, style='Light Grid Accent 1')
data = [('Componente','Tecnologia'),('Frontend','HTML5, CSS3, JavaScript ES6+ (Vanilla)'),
    ('Backend','PHP 8.x con API REST'),('Base de Datos','MySQL 8.0 / MariaDB'),
    ('Mapas','Leaflet.js + OpenStreetMap'),('Exportacion','SheetJS (xlsx) via CDN'),
    ('Servidor','Apache (XAMPP)'),('PWA','Service Worker + Manifest')]
for i,(a,b) in enumerate(data):
    table.rows[i].cells[0].text = a; table.rows[i].cells[1].text = b
spacer()
title('2.2 Estructura de Archivos', level=2)
for f in ['index.html — Punto de entrada principal','src/main.js — Router y logica de navegacion',
    'src/pages/ — Paginas (Home, Login, Trips, Profile, Admin, Purchase)',
    'src/services/ — Servicios (Auth, Data, Map, Excel, Router)',
    'src/components/ — Componentes reutilizables (Navbar, Toast, Icons)',
    'src/styles/ — Hojas de estilo (tokens, global, components, pages)',
    'api/ — Backend PHP para MySQL','sw.js — Service Worker para modo offline']:
    bullet(f)
title('2.3 Almacenamiento de Datos', level=2)
para('El sistema usa un enfoque dual: MySQL para produccion y localStorage como respaldo para modo demo/offline. Las claves principales de almacenamiento son:')
for k in ['premiumbus_users — Registro de usuarios','premiumbus_trips — Rutas y paradas',
    'premiumbus_purchases — Boletos activos','premiumbus_history — Viajes completados',
    'premiumbus_drivers — Conductores registrados']:
    bullet(k)
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 3: INSTALACION PC
# ══════════════════════════════════════════════
title('3. Instalacion en PC (Windows)')
para('La instalacion en PC requiere un servidor web local (XAMPP) para ejecutar el backend PHP y la base de datos MySQL.')
title('3.1 Requisitos Previos', level=2)
bullet('Windows 10 o superior'); bullet('XAMPP instalado (descarga: https://www.apachefriends.org/)')
bullet('50 MB de espacio libre'); bullet('Navegador web moderno (Chrome, Edge, Firefox)')
title('3.2 Pasos de Instalacion', level=2)
for s in ['Descargue e instale XAMPP desde apachefriends.org si aun no lo tiene.',
    'Descargue el proyecto PremiumBus (ZIP de GitHub o ejecute el instalador .exe).',
    'Si usa el instalador: ejecute Instalador_PremiumBus_v5.exe y siga el asistente de 4 pasos.',
    'Si instala manualmente: copie la carpeta PremiumBus dentro de C:\\xampp\\htdocs\\.',
    'Abra XAMPP Control Panel e inicie los servicios Apache y MySQL.',
    'Abra su navegador y visite: http://localhost/PremiumBus/',
    'La primera vez, visite http://localhost/PremiumBus/setup_db.php para crear la base de datos.',
    'Inicie sesion con las credenciales de administrador: admin@premiumbus.com / admin123']:
    numbered(s)
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 4: INSTALACION ANDROID
# ══════════════════════════════════════════════
title('4. Instalacion en Android')
para('PremiumBus es una PWA que se instala directamente desde el navegador Chrome sin necesidad de tienda de aplicaciones.')
title('4.1 Pasos de Instalacion', level=2)
for s in ['Abra Google Chrome en su dispositivo Android.',
    'Navegue a la URL del servidor: http://[IP-del-servidor]/PremiumBus/ (o la URL de produccion).',
    'Espere a que la pagina cargue completamente.',
    'Chrome mostrara automaticamente un banner inferior que dice "Agregar PremiumBus a la pantalla de inicio". Toque "Instalar".',
    'Si no aparece el banner: toque el menu de tres puntos (esquina superior derecha) y seleccione "Instalar aplicacion" o "Agregar a pantalla de inicio".',
    'Confirme la instalacion tocando "Instalar" en el dialogo.',
    'PremiumBus aparecera como un icono en su pantalla de inicio.',
    'Abra la aplicacion desde el icono. Se ejecutara en modo pantalla completa como una app nativa.']:
    numbered(s)
title('4.2 Requisitos', level=2)
bullet('Android 8.0 o superior'); bullet('Google Chrome actualizado'); bullet('Conexion a internet para la primera carga')
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 5: INSTALACION IOS
# ══════════════════════════════════════════════
title('5. Instalacion en iOS (iPhone / iPad)')
para('En dispositivos Apple, la instalacion se realiza a traves de Safari utilizando la funcion "Agregar a inicio".')
title('5.1 Pasos de Instalacion', level=2)
for s in ['Abra Safari en su iPhone o iPad (debe ser Safari, no Chrome).',
    'Navegue a la URL del servidor PremiumBus.',
    'Espere a que la pagina cargue completamente.',
    'Toque el icono de Compartir (cuadrado con flecha hacia arriba) en la barra inferior de Safari.',
    'En el menu que aparece, desplacese hacia abajo y seleccione "Agregar a pantalla de inicio".',
    'Edite el nombre si lo desea (por defecto sera "PremiumBus") y toque "Agregar".',
    'La aplicacion aparecera como un icono en su pantalla de inicio.',
    'Abra la app desde el icono. Se ejecutara en pantalla completa sin barras de Safari.']:
    numbered(s)
title('5.2 Requisitos', level=2)
bullet('iOS 14.0 o superior'); bullet('Safari (obligatorio para PWA en iOS)'); bullet('Conexion a internet para la primera carga')
title('5.3 Nota Importante', level=2)
para('En iOS, las PWA tienen algunas limitaciones: no reciben notificaciones push (hasta iOS 16.4+) y los datos se eliminan si no se usa la app durante varias semanas. Se recomienda abrir la app al menos una vez por semana.', italic=True)
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 6: GUIA USUARIO
# ══════════════════════════════════════════════
title('6. Guia de Uso — Usuario')
title('6.1 Registro e Inicio de Sesion', level=2)
for s in ['Abra la aplicacion PremiumBus.','Toque "Crear Cuenta" en la pantalla de login.',
    'Ingrese su nombre, correo electronico y contrasena.','Toque "Registrarse" para crear su cuenta.',
    'Inicie sesion con su correo y contrasena.']:
    numbered(s)
title('6.2 Consultar Rutas y Mapa', level=2)
for s in ['Desde la pantalla principal, toque "Viajes" en la barra de navegacion inferior.',
    'Vera la lista de rutas disponibles con precios y horarios.',
    'Toque una ruta para ver el mapa con las paradas marcadas.',
    'La ruta se dibuja en color cyan sobre el mapa.',
    'Toque "En Vivo" para activar la simulacion del autobus en tiempo real.',
    'Toque "Mi Ubicacion" para ver su posicion GPS en el mapa.']:
    numbered(s)
title('6.3 Comprar un Boleto', level=2)
for s in ['Seleccione una ruta y toque el boton "Comprar".','Elija su asiento en el diagrama del autobus (los verdes estan disponibles).',
    'Confirme la compra. Se generara un boleto digital con codigo QR.',
    'El boleto aparecera en su perfil con toda la informacion del viaje.']:
    numbered(s)
title('6.4 Mi Perfil', level=2)
para('Desde el icono de Perfil puede: ver sus boletos activos, historial de viajes, cambiar contrasena, actualizar foto de perfil y eliminar su cuenta.')
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 7: GUIA ADMIN
# ══════════════════════════════════════════════
title('7. Guia de Uso — Administrador')
para('Los administradores tienen acceso a funciones avanzadas de gestion del sistema.')
title('7.1 Acceso al Panel de Administracion', level=2)
para('Inicie sesion con una cuenta de administrador (admin@premiumbus.com / admin123). Desde la pantalla principal, toque "Panel Admin" para acceder al panel completo.')
title('7.2 Pestanas del Panel Admin', level=2)
para('El panel tiene 4 pestanas:', bold=True)
bullet('Rutas: Vista de todas las rutas con ocupacion, conductor asignado y datos historicos')
bullet('Conductores: Agregar nuevos conductores, eliminar existentes, ver asignaciones de ruta')
bullet('Usuarios: Lista de todos los usuarios registrados con boton de eliminacion (solo admin principal)')
bullet('Compras: Historial de todas las compras realizadas en el sistema')
title('7.3 Exportar Reporte Excel', level=2)
para('Toque el boton "Exportar Reporte Excel (.xlsx)" en el panel admin. Se generara un archivo con 6 hojas: Compras Activas, Historial, Usuarios, Rutas, Conductores y Resumen Ejecutivo. Este archivo es ideal para auditorias.')
title('7.4 Eliminar Usuarios', level=2)
para('Solo el administrador principal (ID 1) puede eliminar usuarios y otros administradores. Al eliminar un usuario se borran todos sus datos: compras, historial y foto de perfil. Esta accion es irreversible.')
title('7.5 Buscador Avanzado de Usuarios', level=2)
para('En el perfil de administrador, la seccion de usuarios incluye un buscador con algoritmo fuzzy que permite buscar por nombre, correo o ID. Incluye filtros por rol (Todos, Usuarios, Admins) y resalta el texto coincidente.')
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 8: FUNCIONALIDADES v5.0
# ══════════════════════════════════════════════
title('8. Funcionalidades de la Version 5.0')
para('La version 5.0 introduce mejoras significativas en administracion, reportes y experiencia de mapa:')
features = [
    ('Eliminacion de usuarios por admin', 'El administrador principal puede eliminar cualquier usuario o administrador del sistema desde el Panel Admin o el Perfil.'),
    ('Exportacion Excel profesional', 'Genera archivos .xlsx reales con 6 hojas de datos usando la libreria SheetJS, ideal para auditorias y reportes financieros.'),
    ('CRUD de conductores', 'Gestion completa de conductores: agregar con nombre, telefono y licencia; eliminar; asignar a rutas especificas.'),
    ('Mapa con colores neon', 'Rutas dibujadas en cyan neon (#00D4FF) con efecto de brillo. Recorrido en azul neon (#0080FF). Marcadores con efecto glow.'),
    ('Buscador fuzzy avanzado', 'Algoritmo de busqueda por tokens con scoring ponderado, normalizacion unicode (insensible a acentos) y ordenamiento por relevancia.'),
    ('Filtros por rol', 'Pills de filtrado rapido: Todos, Usuarios, Admins con conteo en tiempo real.'),
]
for feat_title, feat_desc in features:
    p = doc.add_paragraph()
    r1 = p.add_run(f'{feat_title}: ')
    r1.bold = True; r1.font.size = Pt(11)
    r2 = p.add_run(feat_desc)
    r2.font.size = Pt(11)
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 9: RECUPERACION BD
# ══════════════════════════════════════════════
title('9. Recuperacion de Base de Datos por Emergencias')
para('Esta seccion describe los procedimientos para proteger y recuperar la base de datos del sistema en caso de fallas, corrupcion de datos o perdida accidental.')
title('9.1 Respaldo Manual de la Base de Datos', level=2)
para('Metodo 1: Usando phpMyAdmin', bold=True)
for s in ['Abra su navegador y visite http://localhost/phpmyadmin/',
    'Seleccione la base de datos "premiumbus_db" en el panel izquierdo.',
    'Haga clic en la pestana "Exportar" en la barra superior.',
    'Seleccione el metodo "Rapido" y formato "SQL".',
    'Haga clic en "Continuar". Se descargara un archivo .sql con todo el respaldo.']:
    numbered(s)
spacer()
para('Metodo 2: Usando la linea de comandos', bold=True)
p = doc.add_paragraph()
r = p.add_run('mysqldump -u root -p premiumbus_db > respaldo_premiumbus.sql')
r.font.name = 'Consolas'; r.font.size = Pt(10)
title('9.2 Restauracion de un Respaldo', level=2)
para('Metodo 1: Usando phpMyAdmin', bold=True)
for s in ['Abra phpMyAdmin (http://localhost/phpmyadmin/).',
    'Si la base de datos existe, seleccionela y vaya a "Operaciones" > "Eliminar base de datos".',
    'Cree una nueva base de datos con el nombre "premiumbus_db".',
    'Seleccione la nueva base de datos y vaya a la pestana "Importar".',
    'Haga clic en "Seleccionar archivo" y elija el archivo .sql de respaldo.',
    'Haga clic en "Continuar". La base de datos sera restaurada completamente.']:
    numbered(s)
spacer()
para('Metodo 2: Usando la linea de comandos', bold=True)
p = doc.add_paragraph()
r = p.add_run('mysql -u root -p premiumbus_db < respaldo_premiumbus.sql')
r.font.name = 'Consolas'; r.font.size = Pt(10)
title('9.3 Respaldo Automatico (Recomendado)', level=2)
para('Para automatizar respaldos diarios, cree un archivo .bat con el siguiente contenido y programelo en el Programador de Tareas de Windows:')
p = doc.add_paragraph()
code = '@echo off\nset FECHA=%date:~6,4%-%date:~3,2%-%date:~0,2%\n"C:\\xampp\\mysql\\bin\\mysqldump.exe" -u root premiumbus_db > "C:\\Respaldos\\premiumbus_%FECHA%.sql"'
r = p.add_run(code)
r.font.name = 'Consolas'; r.font.size = Pt(9)
title('9.4 Recuperacion de Datos de localStorage', level=2)
para('Si el sistema opera en modo demo (sin MySQL), los datos se almacenan en el navegador. Para respaldar:')
for s in ['Abra la aplicacion en Chrome.','Presione F12 para abrir las herramientas de desarrollador.',
    'Vaya a la pestana "Application" > "Local Storage" > URL del sitio.',
    'Copie manualmente los valores de las claves premiumbus_users, premiumbus_trips, premiumbus_purchases, premiumbus_history y premiumbus_drivers.',
    'Guarde los valores en un archivo JSON como respaldo.']:
    numbered(s)
title('9.5 Plan de Contingencia', level=2)
bullet('Frecuencia de respaldo recomendada: diario (automatico) + semanal (manual)')
bullet('Ubicacion de respaldos: disco externo o servicio en la nube (Google Drive, OneDrive)')
bullet('Tiempo estimado de recuperacion: 5-10 minutos con respaldo reciente')
bullet('Contacto de soporte: administrador del sistema (admin@premiumbus.com)')
doc.add_page_break()

# ══════════════════════════════════════════════
# CAP 10: FAQ
# ══════════════════════════════════════════════
title('10. Preguntas Frecuentes')
faqs = [
    ('No puedo iniciar sesion', 'Verifique que su correo y contrasena sean correctos. Si olvido su contrasena, contacte al administrador para que la restablezca.'),
    ('La aplicacion no carga el mapa', 'Verifique su conexion a internet. El mapa requiere acceso a los servidores de OpenStreetMap para cargar los tiles.'),
    ('No puedo comprar boletos', 'Asegurese de haber iniciado sesion como usuario (no como administrador). Los administradores no pueden comprar boletos.'),
    ('Como cambio mi contrasena', 'Vaya a Perfil > Configuracion > Cambiar Contrasena. Ingrese su contrasena actual y la nueva.'),
    ('Como exporto datos para auditoria', 'Desde el Panel Admin o el Perfil de administrador, toque el boton "Exportar Reporte Excel". Se descargara un archivo .xlsx.'),
    ('La app se ve rara en mi celular', 'Asegurese de usar un navegador actualizado. En iOS debe usar Safari. En Android use Chrome.'),
    ('Como agrego un conductor nuevo', 'Panel Admin > pestana Conductores > boton "Agregar Conductor". Complete nombre, telefono y licencia.'),
    ('Que hago si borre un usuario por error', 'La eliminacion es irreversible. Restaure la base de datos desde un respaldo reciente (ver Capitulo 9).'),
]
for q, a in faqs:
    p = doc.add_paragraph()
    r1 = p.add_run(f'P: {q}')
    r1.bold = True; r1.font.size = Pt(11)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(f'R: {a}')
    r2.font.size = Pt(11)
    p2.paragraph_format.space_after = Pt(12)

spacer()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('— Fin del Manual de Usuario —')
r.font.size = Pt(14); r.bold = True; r.font.color.rgb = RGBColor(26, 58, 107)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('PremiumBus v5.0 | San Luis Potosi, Mexico | Mayo 2026')
r2.font.size = Pt(10); r2.font.color.rgb = RGBColor(150, 150, 150)

# Guardar
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Manual_Usuario_PremiumBus_v5.docx')
doc.save(out)
print(f'\n[OK] Manual generado: {out}')
print(f'     Paginas estimadas: 13+')
os.startfile(out)
