"""
Genera la presentacion PowerPoint de pruebas PremiumBus.
Requiere: pip install python-pptx
Ejecutar: python generar_pptx.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

DARK_BG = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1A, 0x3A, 0x6B)
BLUE2 = RGBColor(0x2B, 0x5E, 0xA7)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
RED = RGBColor(0xEF, 0x44, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
LGRAY = RGBColor(0xE2, 0xE8, 0xF0)
CARD = RGBColor(0x1E, 0x29, 0x3B)

def set_bg(slide, color=DARK_BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_text(slide, text, x, y, w, h, size=12, color=LGRAY, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = align
    # Handle newlines
    if '\n' in text:
        tf.clear()
        for i, line in enumerate(text.split('\n')):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.color.rgb = color
            p.font.bold = bold
            p.alignment = align

def add_card(slide, x, y, w, h):
    shape = slide.shapes.add_shape(
        5, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.fill.background()

# ===== SLIDE 1: Portada =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BLUE)
add_text(s, 'Plan de Pruebas de Software', 0.5, 1.5, 12.3, 1, 40, WHITE, True, PP_ALIGN.CENTER)
add_text(s, 'PremiumBus — Metodologia Pressman (Cap. 17-20)', 0.5, 2.8, 12.3, 0.5, 20, LGRAY, False, PP_ALIGN.CENTER)
add_text(s, '10 pruebas ejecutadas  |  10 aprobadas  |  100% cobertura', 0.5, 3.6, 12.3, 0.5, 14, GRAY, False, PP_ALIGN.CENTER)
add_text(s, 'Mayo 2026', 0.5, 5.5, 12.3, 0.3, 13, GRAY, False, PP_ALIGN.CENTER)

# ===== SLIDE 2: Autenticacion =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 'Modulo de Autenticacion', 0.5, 0.2, 12, 0.6, 28, WHITE, True)
add_card(s, 0.4, 1, 6, 5.8)
add_text(s, 'CP-AUTH-01: Login y Retencion de Token', 0.6, 1.1, 5.5, 0.4, 14, BLUE2, True)
add_text(s, 'Se ingresaron credenciales validas.\nSe verifico acceso al Home.\nSe simulo cierre y reapertura de la app.\n\nEl token se mantiene en localStorage.\nLa sesion persiste sin pedir credenciales.\n\nResultado: APROBADA', 0.6, 1.6, 5.5, 4.5, 12, LGRAY)
add_card(s, 6.8, 1, 6, 5.8)
add_text(s, 'CP-AUTH-02: Prevencion XSS/SQL', 7, 1.1, 5.5, 0.4, 14, BLUE2, True)
add_text(s, 'Se inyectaron 3 vectores de ataque:\n- <script>alert(1)</script>\n- \' OR 1=1 --\n- <img onerror=alert(1)>\n\nTodos bloqueados por validacion regex\nantes de enviar peticion de red.\n(Early Return Pattern)\n\nResultado: APROBADA', 7, 1.6, 5.5, 4.5, 12, LGRAY)

# ===== SLIDE 3: Compras =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 'Modulo de Compras — Valores Limite', 0.5, 0.2, 12, 0.6, 28, WHITE, True)
add_card(s, 0.4, 1, 6, 5.8)
add_text(s, 'CP-COMPRA-01: Saldo Exacto $13.50', 0.6, 1.1, 5.5, 0.4, 14, GREEN, True)
add_text(s, 'Saldo del usuario = $13.50\nPrecio del boleto = $13.50\n\nSe selecciono ruta, asiento y se confirmo.\nLa transaccion fue aprobada.\nCodigo QR emitido.\n\nSaldo final: $0.00 exacto.\n\nResultado: APROBADA', 0.6, 1.6, 5.5, 4.5, 12, LGRAY)
add_card(s, 6.8, 1, 6, 5.8)
add_text(s, 'CP-COMPRA-02: Saldo Insuficiente $13.49', 7, 1.1, 5.5, 0.4, 14, RED, True)
add_text(s, 'Saldo = $13.49 (1 centavo bajo el limite)\nPrecio = $13.50\nDiferencia = -$0.01\n\nToast rojo: "Fondo insuficiente"\nTransaccion abortada.\n\nSaldo preservado sin deduccion.\n\nResultado: APROBADA', 7, 1.6, 5.5, 4.5, 12, LGRAY)

# ===== SLIDE 4: Roles =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 'Modulo de Roles y Acceso', 0.5, 0.2, 12, 0.6, 28, WHITE, True)
add_card(s, 0.4, 1, 6, 5.8)
add_text(s, 'CP-ROL-01: Restriccion Vistas Admin', 0.6, 1.1, 5.5, 0.4, 14, BLUE2, True)
add_text(s, 'Con cuenta Admin logueada se verifico:\n\nPanel Conductores -> VISIBLE\nEstadisticas -> VISIBLE\nBoton Comprar -> OCULTO\nQuick Action Comprar -> OCULTO\n\nRenderizado condicional correcto.\n\nResultado: APROBADA', 0.6, 1.6, 5.5, 4.5, 12, LGRAY)
add_card(s, 6.8, 1, 6, 5.8)
add_text(s, 'CP-ROL-02: Directorio de Usuarios', 7, 1.1, 5.5, 0.4, 14, BLUE2, True)
add_text(s, 'Se navego al directorio como Admin.\n\nLista completa cargada con\nnombre, correo y rol de cada usuario.\n\nSin errores de composite indices\nen consola de Firebase.\n\nResultado: APROBADA', 7, 1.6, 5.5, 4.5, 12, LGRAY)

# ===== SLIDE 5: Perfil =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 'Modulo de Perfil y Rendimiento', 0.5, 0.2, 12, 0.6, 28, WHITE, True)
add_card(s, 0.4, 1, 12.4, 5.8)
add_text(s, 'CP-PERF-01: Carga de Fotografia de Perfil', 0.6, 1.1, 12, 0.4, 14, BLUE2, True)
add_text(s, 'Se abrio el selector de imagen.\nSe cargo foto JPG valida (1.2 MB).\n\nDurante la subida se mostro barra de progreso\nasincrona (0% a 100%).\n\nAl completar, la miniatura del avatar se actualizo\nen tiempo real sin recargar la pantalla.\n\nIntegracion con Firebase Storage validada.\n\nResultado: APROBADA', 0.6, 1.6, 12, 4.5, 13, LGRAY)

# ===== SLIDE 6: Geolocalizacion =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 'Modulo de Geolocalizacion', 0.5, 0.2, 12, 0.6, 28, WHITE, True)
add_card(s, 0.4, 1, 6, 5.8)
add_text(s, 'CP-GEO-01: Tracking en Tiempo Real', 0.6, 1.1, 5.5, 0.4, 14, BLUE2, True)
add_text(s, 'Con GPS activado se abrio el mapa.\n\nEl marcador del bus se desplazo\nfluidamente sin saltos.\n\nCoordenadas y velocidad actualizadas\nen tiempo real.\n\nPoller optimizado.\n\nResultado: APROBADA', 0.6, 1.6, 5.5, 4.5, 12, LGRAY)
add_card(s, 6.8, 1, 6, 5.8)
add_text(s, 'CP-GEO-02: Corte de GPS', 7, 1.1, 5.5, 0.4, 14, RED, True)
add_text(s, 'Se desactivo GPS durante tracking.\n\nLa app detecto el cambio de estado,\ncongelo la actualizacion del mapa\ny mostro Toast:\n"Se perdio senal GPS"\n\nSin crash. App estable.\n\nResultado: APROBADA', 7, 1.6, 5.5, 4.5, 12, LGRAY)

# ===== SLIDE 7: UI + Videos =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s)
add_text(s, 'Validacion UI/UX — Cap. 20 WebApp', 0.5, 0.2, 12, 0.6, 28, WHITE, True)
add_card(s, 0.4, 1, 6, 5.8)
add_text(s, 'CP-UI-01: Tokens de Diseno Atomico', 0.6, 1.1, 5.5, 0.4, 14, BLUE2, True)
add_text(s, 'Tokens verificados:\n- --color-primary-700: #1a3a6b\n- --color-accent-green: #34c759\n- --font-family-base: Inter\n- --radius-md: 16px\n\nComponentes aislados:\nToast.js / Navbar.js / Icons.js\n\nToast <= 85% ancho pantalla\nResponsive max-width: 480px\n\nResultado: APROBADA', 0.6, 1.6, 5.5, 4.5, 12, LGRAY)
add_card(s, 6.8, 1, 6, 5.8)
add_text(s, 'Videos de Evidencia', 7, 1.1, 5.5, 0.4, 14, GREEN, True)
add_text(s, 'Cada prueba fue grabada en video:\n\n1. CP-AUTH-01_Login_Exitoso.mp4\n2. CP-AUTH-02_Prevencion_XSS.mp4\n3. CP-PERF-01_Foto_Perfil.mp4\n4. CP-COMPRA-01_Saldo_Exacto.mp4\n5. CP-COMPRA-02_Saldo_Insuficiente.mp4\n6. CP-ROL-01_Restriccion_Admin.mp4\n7. CP-ROL-02_Directorio_Usuarios.mp4\n8. CP-GEO-01_GPS_Tracking.mp4\n9. CP-GEO-02_Corte_GPS.mp4\n10. CP-UI-01_Tokens_Diseno.mp4', 7, 1.6, 5.5, 4.5, 12, LGRAY)

# ===== SLIDE 8: Conclusiones =====
s = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(s, BLUE)
add_text(s, 'Conclusiones', 0.5, 0.3, 12.3, 0.6, 30, WHITE, True, PP_ALIGN.CENTER)
add_card(s, 0.8, 1.2, 11.7, 5.2)
add_text(s, '- 10 de 10 pruebas aprobadas (100%)\n- Cero incidentes de Severidad 1\n- Cobertura completa: Auth, Compras, Roles, Perfil, GPS, UI\n\n- Pruebas Criticas (AUTH, COMPRA, ROL): 100%\n- Sin caidas, fugas de datos ni corrupcion de saldos\n- Codigo cumple SOLID, Early Return y SoC\n- Tokens de diseno atomico validados\n\nEl sistema PremiumBus cumple los criterios de aceptacion\nde Pressman y es apto para produccion.', 1.2, 1.6, 10.8, 4.4, 15, LGRAY)
add_text(s, 'PremiumBus — Mayo 2026', 0.5, 6.8, 12.3, 0.3, 12, GRAY, False, PP_ALIGN.CENTER)

# Guardar
out = os.path.join(os.path.dirname(__file__), 'Presentacion_Pruebas_PremiumBus.pptx')
prs.save(out)
print(f'Presentacion guardada en: {out}')
