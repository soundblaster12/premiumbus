import os
import glob
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
import PyPDF2

def create_element(name): 
    return OxmlElement(name)

def create_attribute(element, name, value): 
    element.set(qn(name), value)

def add_page_number(run):
    fldChar1 = create_element('w:fldChar')
    create_attribute(fldChar1, 'w:fldCharType', 'begin')
    instrText = create_element('w:instrText')
    create_attribute(instrText, 'xml:space', 'preserve')
    instrText.text = "PAGE"
    fldChar2 = create_element('w:fldChar')
    create_attribute(fldChar2, 'w:fldCharType', 'separate')
    fldChar3 = create_element('w:fldChar')
    create_attribute(fldChar3, 'w:fldCharType', 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

doc = Document()

# Margen Oficial TecNM
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    # Encabezado y Pie de página
    add_page_number(section.footer.paragraphs[0].add_run())
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

# Estilo de Texto Normal
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
style.paragraph_format.first_line_indent = Cm(1.25)

def add_p_clean(text, bold=False, italic=False, center=False, level=None):
    text_str = text.strip()
    if not text_str:
        return None
    if level:
        h = doc.add_heading(text_str, level=level)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(6)
        return h
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    
    run = p.add_run(text_str)
    run.bold = bold
    run.italic = italic
    return p

def extract_pdf_text(folder_path):
    text = ""
    for root, dirs, files in os.walk(folder_path):
        for file in files:
            if file.endswith('.pdf'):
                try:
                    with open(os.path.join(root, file), 'rb') as f:
                        reader = PyPDF2.PdfReader(f)
                        for page in reader.pages:
                            t = page.extract_text()
                            if t:
                                text += t + "\n"
                except Exception as e:
                    print(f"Error al leer PDF {file}: {e}")
    return text

def add_pdf_text_to_docx(folder_path, title_name):
    add_p_clean(title_name, level=2)
    text = extract_pdf_text(folder_path)
    if not text.strip():
        add_p_clean("[No se encontraron archivos PDF o contenido legible en esta sección.]", italic=True)
        return
    
    lines = text.split('\n')
    current_paragraph = ""
    for line in lines:
        line_str = line.strip()
        if not line_str:
            if current_paragraph:
                add_p_clean(current_paragraph)
                current_paragraph = ""
            continue
        
        # Filtros de ruido típicos de PDFs de tareas
        if "INGENIERÍA DE SOFTWARE" in line_str.upper() or "CORONADO HERNÁNDEZ DAVID" in line_str.upper() or "PROFESOR" in line_str.upper() or "PÁGINA" in line_str.upper():
            continue
            
        # Reconstruir párrafos
        # Si termina en punto, es una línea muy corta, o empieza por viñeta, cerramos el párrafo actual.
        if line_str.endswith('.') or len(line_str) < 50 or line_str.startswith('•') or line_str.startswith('-') or line_str.startswith('*') or (line_str[0].isdigit() and line_str.endswith(':')):
            if current_paragraph:
                current_paragraph += " " + line_str
                add_p_clean(current_paragraph)
                current_paragraph = ""
            else:
                add_p_clean(line_str)
        else:
            if current_paragraph:
                current_paragraph += " " + line_str
            else:
                current_paragraph = line_str
    if current_paragraph:
        add_p_clean(current_paragraph)

def append_docx_clean(src_path):
    if not os.path.exists(src_path):
        print(f"Saltando {src_path} porque no existe.")
        return
    try:
        src_doc = Document(src_path)
        for element in src_doc.element.body:
            if element.tag.endswith('p'):
                p = Paragraph(element, src_doc)
                p_text = p.text.strip()
                if p_text:
                    # Omitir metadatos repetitivos de las portadas originales
                    if p_text in [
                        "Materia:", "Ingeniería de software", "8:00 - 9:00 AM", 
                        "Maestro: Oscar Abundio Juárez Romero", "Fecha de entrega: 25 de mayo del 2026", 
                        "Integrantes:", "Coronado Hernández David", "Gonzales Castillo Gael Essau", 
                        "Moreno Flores Braulio Edgardo", "Reyna Hernández Yaotl Isaias", 
                        "Ing. Sistemas Computacionales", "Fecha de entrega: 25 de mayo del 2026"
                    ]:
                        continue
                    if "Tarea 4.4 Documentación" in p_text or "Manual de usuario - PremiumBus" in p_text:
                        continue
                    
                    is_heading = False
                    level = 1
                    if p.style.name.startswith('Heading'):
                        is_heading = True
                        try:
                            level = int(p.style.name.replace('Heading', ''))
                        except:
                            level = 2
                    elif p_text.isupper() and len(p_text) < 60:
                        is_heading = True
                        level = 2
                    elif p_text.startswith(('1. ', '2. ', '3. ', '4. ', '5. ', '6. ', '7. ', '8. ', '9. ', 'Capítulo')):
                        is_heading = True
                        level = 2
                    
                    if is_heading:
                        h = doc.add_heading(p_text, level=min(level, 3))
                        h.paragraph_format.first_line_indent = Cm(0)
                        h.paragraph_format.space_before = Pt(12)
                        h.paragraph_format.space_after = Pt(6)
                    else:
                        dp = doc.add_paragraph()
                        dp.paragraph_format.first_line_indent = Cm(1.25)
                        dp.paragraph_format.line_spacing = 1.5
                        dp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        for run in p.runs:
                            dr = dp.add_run(run.text)
                            dr.bold = run.bold
                            dr.italic = run.italic
                            dr.underline = run.underline
            elif element.tag.endswith('tbl'):
                table = Table(element, src_doc)
                dt = doc.add_table(rows=len(table.rows), cols=len(table.columns))
                dt.style = 'Table Grid'
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        dest_cell = dt.cell(r_idx, c_idx)
                        # Limpiar párrafos por defecto
                        for cp in list(dest_cell.paragraphs):
                            dest_cell._element.remove(cp._element)
                        for p in cell.paragraphs:
                            dp = dest_cell.add_paragraph()
                            dp.paragraph_format.line_spacing = 1.15
                            dp.paragraph_format.space_after = Pt(2)
                            for run in p.runs:
                                dr = dp.add_run(run.text)
                                dr.bold = run.bold
                                dr.italic = run.italic
                                dr.underline = run.underline
    except Exception as e:
        print(f"Error al acoplar docx {src_path}: {e}")

# ==============================================================================
# --- PORTADA DE DOCUMENTACIÓN DE SOFTWARE ---
# ==============================================================================
doc.add_paragraph('\n\n')
add_p_clean('INSTITUTO TECNOLÓGICO', bold=True, center=True)
add_p_clean('(TECNM)', bold=True, center=True)
doc.add_paragraph('\n\n')
add_p_clean('DOCUMENTACIÓN DE SISTEMA Y ESPECIFICACIÓN TÉCNICA', bold=True, center=True)
add_p_clean('PremiumBus - Servicio on-line y móvil para gestión logística y boletaje', center=True)
add_p_clean('de transporte terrestre en una Institución de Educación Superior', center=True)
doc.add_paragraph('\n\n')
add_p_clean('Materia:', bold=True, center=True)
add_p_clean('Ingeniería de Software', center=True)
doc.add_paragraph('\n')
add_p_clean('Maestro:', bold=True, center=True)
add_p_clean('Oscar Abundio Juárez Romero', center=True)
doc.add_paragraph('\n')
add_p_clean('Presentan:', bold=True, center=True)
add_p_clean('Coronado Hernández David', center=True)
add_p_clean('Gonzales Castillo Gael Essau', center=True)
add_p_clean('Moreno Flores Braulio Edgardo', center=True)
add_p_clean('Reyna Hernández Yaotl Isaias', center=True)
doc.add_paragraph('\n\n\n')
add_p_clean('Fecha: Mayo de 2026', center=True)
doc.add_page_break()

# ==============================================================================
# --- RESUMEN ---
# ==============================================================================
add_p_clean('RESUMEN DEL SISTEMA', level=1)
add_p_clean('El presente documento constituye la especificación de requisitos y documentación técnica detallada del sistema PremiumBus, una plataforma móvil y web integral para la gestión logística y venta de boletos de transporte para comunidades universitarias. Diseñado para simplificar, optimizar y automatizar el proceso de boletaje, PremiumBus permite tanto a los usuarios generales como a los administradores del sistema gestionar viajes, reservar y bloquear asientos en tiempo real, visualizar rutas con mapas interactivos y auditar ventas de forma centralizada.')
add_p_clean('La especificación se estructura formalmente siguiendo los estándares internacionales de la Ingeniería de Software como ISO/IEC/IEEE 29148:2011 y el estándar IEEE 830-1998, asegurando un ciclo de vida limpio y escalable para la aplicación. El documento fusiona la planeación inicial de requerimientos, análisis de factibilidad, casos de uso, topología de la base de datos, matrices de pruebas de software y el manual de usuario, proporcionando una base documental sólida libre de deuda técnica.')
doc.add_page_break()

# ==============================================================================
# --- CAPÍTULO 1. INTRODUCCIÓN Y ALCANCE (BASADO EN SRS) ---
# ==============================================================================
add_p_clean('CAPÍTULO 1. INTRODUCCIÓN AL SISTEMA', level=1)
add_p_clean('1.1 Propósito de la Documentación', level=2)
add_p_clean('La presente documentación técnica representa el activo intangible más valioso de PremiumBus. Su propósito central es mitigar la deuda técnica y erradicar el riesgo operativo que asume una organización cuando el conocimiento crítico del sistema reside exclusivamente en la memoria de los desarrolladores originales. Al compilar los esquemas de bases de datos, especificaciones de requerimientos, decisiones arquitectónicas y matrices de validación, este documento se convierte en la única fuente de verdad (Single Source of Truth), asegurando una inducción rápida para futuros ingenieros de software, claridad técnica para auditorías y un mapa claro para expansiones a plataformas de nube.')

add_p_clean('1.2 Alcance del Sistema', level=2)
add_p_clean('PremiumBus es una solución integral que abarca desde la visualización de rutas interactivas hasta la asignación lógica de asientos y la venta de boletos. El alcance de la plataforma incluye la gestión dinámica del inventario de autobuses, la orquestación comercial del boleto (búsqueda, bloqueo por concurrencia y validación mediante códigos QR), un panel de administración centralizado y la visualización de recorridos mediante la API de Google Maps.')

add_p_clean('1.3 Definiciones y Acrónimos', level=2)
add_p_clean('• SRS / ERS: Especificación de Requerimientos de Software.')
add_p_clean('• UI: Interfaz de Usuario.')
add_p_clean('• BD / SGBD: Base de Datos / Sistema Gestor de Base de Datos Relacional (MySQL).')
add_p_clean('• API: Interfaz de Programación de Aplicaciones (Application Programming Interface).')
add_p_clean('• CRUD: Operaciones básicas de datos (Crear, Leer, Actualizar, Borrar).')

add_p_clean('1.4 Factibilidad y Análisis de Viabilidad (Unidad 1)', level=2)
add_pdf_text_to_docx('C:/Users/david/OneDrive/Documentos/PremiumBus/U1', "1.4.1 Fusión de Análisis de Requisitos U1")
doc.add_page_break()

# ==============================================================================
# --- CAPÍTULO 2. ESPECIFICACIÓN DE REQUISITOS ---
# ==============================================================================
add_p_clean('CAPÍTULO 2. ESPECIFICACIÓN DE REQUISITOS', level=1)
add_p_clean('2.1 Requerimientos Funcionales principales', level=2)
add_p_clean('• RF1 (Registro de Usuario): Permite la creación de nuevas cuentas de usuario con validación de unicidad de correo.')
add_p_clean('• RF2 (Inicio de Sesión): Autenticación segura de usuarios mediante credenciales registradas.')
add_p_clean('• RF3 (Consulta de Viajes): Despliegue de viajes disponibles indicando ruta, fecha, costo y asientos libres.')
add_p_clean('• RF4 (Compra de Boletos): Registro lógico de boletos en la BD con asignación de butaca específica.')
add_p_clean('• RF5 (Navegación e Interfaces): Transiciones fluidas entre pantallas en dispositivos móviles.')
add_p_clean('• RF6 (Acceso Administrativo): Credenciales especiales para acceder a la consola de auditoría y reportes.')

add_p_clean('2.2 Requerimientos No Funcionales principales', level=2)
add_p_clean('• RNF1 (Seguridad): Encriptación criptográfica de contraseñas de los usuarios y validación para evitar inyección SQL.')
add_p_clean('• RNF2 (Rendimiento): Tiempo de respuesta del backend menor a 3 segundos en consultas complejas.')
add_p_clean('• RNF3 (Usabilidad): Interfaz altamente intuitiva y responsiva diseñada para un uso ágil.')
add_p_clean('• RNF4 (Disponibilidad): Operación garantizada del servidor de base de datos MySQL 24/7.')

add_p_clean('2.3 Tablas Descriptivas de Casos de Uso del Sistema', level=2)
casos_uso = [
    "Registro de Usuario", "Login de Sistema", "Búsqueda de Viajes Activos", "Selección de Asiento", 
    "Bloqueo Temporal de Asiento", "Procesamiento de Pago", "Generación de Boleto QR", 
    "Cancelación de Compra", "Recuperación de Contraseña", "Edición de Perfil", 
    "Creación de Nueva Ruta", "Asignación de Vehículo", "Alta de Choferes", 
    "Reporte de Ventas Diarias", "Reporte de Ocupación por Viaje"
]

for idx, cu_nombre in enumerate(casos_uso, 1):
    add_p_clean(f'Tabla {idx}. Caso de Uso – {cu_nombre}', bold=True, center=True)
    table = doc.add_table(rows=11, cols=2)
    table.style = 'Table Grid'
    
    # Rellenar tabla
    table.cell(0, 0).text = "Escenario:"
    table.cell(0, 1).text = f"MÓDULO OPERATIVO M-{idx}"
    table.cell(1, 0).text = "Caso de Uso:"
    table.cell(1, 1).text = cu_nombre
    table.cell(2, 0).text = "Actor Principal:"
    table.cell(2, 1).text = "Administrador" if idx > 10 else "Usuario Registrado"
    table.cell(3, 0).text = "Propósito:"
    table.cell(3, 1).text = f"Ejecutar y validar lógicamente la acción de {cu_nombre.lower()} dentro del entorno PremiumBus."
    table.cell(4, 0).text = "Resumen:"
    table.cell(4, 1).text = f"El actor inicia el flujo para {cu_nombre.lower()}, interactuando con la interfaz móvil, la cual valida e impacta de forma segura los registros en MySQL."
    table.cell(5, 0).text = "Tipo:"
    table.cell(5, 1).text = "Primario / Esencial."
    
    # Fusión de filas para encabezados internos
    table.cell(6, 0).merge(table.cell(6, 1)).text = "Curso Normal de Eventos"
    table.cell(6, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.cell(6, 0).paragraphs[0].runs[0].font.bold = True
    
    table.cell(7, 0).text = "Acciones del Actor"
    table.cell(7, 1).text = "Respuesta del Sistema"
    table.cell(7, 0).paragraphs[0].runs[0].font.bold = True
    table.cell(7, 1).paragraphs[0].runs[0].font.bold = True
    
    table.cell(8, 0).text = "1. El actor presiona la opción correspondiente en la UI e ingresa los datos solicitados."
    table.cell(8, 1).text = "2. El sistema valida las entradas de datos, procesa la solicitud mediante la API y actualiza la base de datos MySQL."
    
    table.cell(9, 0).merge(table.cell(9, 1)).text = "Cursos Alternos y Excepciones"
    table.cell(9, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.cell(9, 0).paragraphs[0].runs[0].font.bold = True
    
    table.cell(10, 0).merge(table.cell(10, 1)).text = "En caso de falla de red o datos inválidos, la aplicación móviles cancela la transacción local, emite un mensaje de error legible al usuario y revierte los bloqueos."
    
    # Ajustar estilos de párrafos en la tabla
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.first_line_indent = Cm(0)
    
    doc.add_paragraph('\n')

add_pdf_text_to_docx('C:/Users/david/OneDrive/Documentos/PremiumBus/U2', "2.4 Fusión de Diagramas e Interfaces de Requisitos U2")
doc.add_page_break()

# ==============================================================================
# --- CAPÍTULO 3. ARQUITECTURA Y DISEÑO DEL SISTEMA ---
# ==============================================================================
add_p_clean('CAPÍTULO 3. ARQUITECTURA Y DISEÑO DEL SISTEMA', level=1)
add_p_clean('3.1 Diseño de Interfaces y Modelos Visuales', level=2)
add_p_clean('El diseño visual y de interacción de PremiumBus fue desarrollado para garantizar máxima eficiencia logística en pantallas móviles y facilidad de uso. A continuación se muestra la maqueta o modelo visual de la interfaz de usuario implementada para el boletaje y visualización de viajes:')

try:
    doc.add_picture('C:/Users/david/OneDrive/Documentos/PremiumBus/diseño.jpg', width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_p_clean('Figura 1. Modelo de Interfaz Visual del Sistema PremiumBus. (FUENTE: Elaboración propia).', center=True, italic=True)
except Exception as e:
    add_p_clean(f"[La imagen de interfaces diseño.jpg no pudo inyectarse: {e}]", italic=True)

add_p_clean('3.2 Diccionario de Datos del Sistema', level=2)
add_p_clean('El modelo relacional de base de datos de PremiumBus ha sido diseñado bajo la tercera forma normal (3FN) para evitar anomalías en inserciones, actualizaciones y borrados. A continuación se desglosa el diccionario de datos de las tablas esenciales del SGBD:')

tablas_db = [
    ('usuarios', [
        ('id', 'INT', '✔', '✔', 'Clave primaria autoincremental única de cada usuario.'),
        ('nombre', 'VARCHAR(100)', '', '✔', 'Nombre completo de la persona registrada.'),
        ('correo', 'VARCHAR(150)', '', '✔', 'Correo electrónico corporativo / Login ID único.'),
        ('password', 'VARCHAR(255)', '', '✔', 'Hash criptográfico seguro de la contraseña.'),
        ('rol', 'VARCHAR(50)', '', '✔', 'Rol de acceso dentro del sistema: user o admin.')
    ]),
    ('viajes', [
        ('id', 'INT', '✔', '✔', 'Clave primaria autoincremental del viaje programado.'),
        ('ruta', 'VARCHAR(100)', '', '✔', 'Nombre geográfico de la ruta asociada al viaje.'),
        ('fecha', 'DATE', '', '✔', 'Fecha de salida programada para el transporte.'),
        ('precio', 'DECIMAL(10,2)', '', '✔', 'Costo por boleto individual en moneda nacional.'),
        ('asientos_libres', 'INT', '', '✔', 'Número de butacas disponibles para venta.')
    ]),
    ('compras', [
        ('id', 'INT', '✔', '✔', 'Clave primaria del boleto emitido y comprado.'),
        ('usuario_id', 'INT', '', '✔', 'Llave foránea que asocia la compra con un usuario.'),
        ('viaje_id', 'INT', '', '✔', 'Llave foránea que enlaza la compra a un viaje activo.'),
        ('asiento', 'INT', '', '✔', 'Número de asiento reservado por el comprador.'),
        ('fecha_compra', 'TIMESTAMP', '', '✔', 'Registro temporal automatizado del momento del pago.')
    ])
]

for t_idx, (t_name, cols) in enumerate(tablas_db, 1):
    add_p_clean(f'Tabla 20.{t_idx}. Diccionario de Datos de la Tabla: {t_name}', bold=True, center=True)
    t_dic = doc.add_table(rows=len(cols)+1, cols=5)
    t_dic.style = 'Table Grid'
    
    # Encabezado
    h = t_dic.rows[0].cells
    h[0].text, h[1].text, h[2].text, h[3].text, h[4].text = ('Campo', 'Tipo', 'PK', 'NN', 'Descripción')
    for cell in h:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    for r_idx, row_data in enumerate(cols, 1):
        for c_idx in range(5):
            t_dic.cell(r_idx, c_idx).text = row_data[c_idx]
            
    # Ajustar estilos
    for row in t_dic.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.first_line_indent = Cm(0)
    doc.add_paragraph('\n')

add_pdf_text_to_docx('C:/Users/david/OneDrive/Documentos/PremiumBus/U3', "3.3 Fusión de Selección de Lenguaje y Despliegue U3")
doc.add_page_break()

# ==============================================================================
# --- CAPÍTULO 4. PRUEBAS DE SOFTWARE (FUSIÓN U4 Y SRS) ---
# ==============================================================================
add_p_clean('CAPÍTULO 4. ELABORACIÓN Y APLICACIÓN DE PRUEBAS', level=1)
add_pdf_text_to_docx('C:/Users/david/OneDrive/Documentos/PremiumBus/U4', "4.1 Fusión de Matrices de Cobertura y Actas de Pruebas U4")
doc.add_page_break()

# ==============================================================================
# --- ACOPLAR TAREA 4.4 DOCUMENTACIÓN COMPLETA (359 párrafos, 9 tablas) ---
# ==============================================================================
add_p_clean('CAPÍTULO 5. ESPECIFICACIÓN Y DISEÑO COMPILADO DEL SISTEMA', level=1)
add_p_clean('La siguiente sección incorpora de forma directa los resultados completos de la especificación, contrato, planes de prueba de caja negra/blanca, benchmarking de endpoints e hitos del proyecto compilados por el equipo de ingeniería de software:', italic=True)
append_docx_clean('C:/Users/david/OneDrive/Documentos/PremiumBus/Tarea 4.4 Documentación del sistema - INGENIERIA DE SOFTWARE U4.docx')
doc.add_page_break()

# ==============================================================================
# --- ACOPLAR MANUAL DE USUARIO INTEGRADO (212 párrafos, 1 tabla) ---
# ==============================================================================
add_p_clean('CAPÍTULO 6. MANUAL DE USUARIO INTEGRADO DEL SISTEMA', level=1)
add_p_clean('Para garantizar un correcto onboarding del software y comprensión operativa de las pantallas desarrolladas, se anexa a continuación el Manual de Usuario integral del sistema:', italic=True)
append_docx_clean('C:/Users/david/OneDrive/Documentos/PremiumBus/Manual_Usuario_PremiumBus_v5.docx')

# Guardar Documento Final Limpio
output_path = 'C:/Users/david/OneDrive/Documentos/PremiumBus/Documentacion_PremiumBus_Software.docx'
doc.save(output_path)
print("¡Documentación oficial sin código generada exitosamente!")
