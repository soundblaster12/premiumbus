import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

def apply_theme(slide, is_title=False):
    # Set background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(30, 40, 50) # Dark modern blue/gray
    
    # Title
    if slide.shapes.title:
        title = slide.shapes.title
        for p in title.text_frame.paragraphs:
            p.font.color.rgb = RGBColor(0, 188, 212) # Teal
            p.font.bold = True
            p.font.name = "Arial"
            if is_title:
                p.font.size = Pt(44)
    
    # Body
    for shape in slide.shapes:
        if not shape.has_text_frame or shape == slide.shapes.title:
            continue
        for p in shape.text_frame.paragraphs:
            p.font.color.rgb = RGBColor(230, 230, 230)
            p.font.name = "Arial"

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "GymSystem Móvil"
slide.placeholders[1].text = "Sistema de Gestión de Gimnasio\nTópicos Avanzados de Programación"
apply_theme(slide, is_title=True)
slide.placeholders[1].text_frame.paragraphs[0].font.color.rgb = RGBColor(200, 200, 200)
if len(slide.placeholders[1].text_frame.paragraphs) > 1:
    slide.placeholders[1].text_frame.paragraphs[1].font.color.rgb = RGBColor(150, 150, 150)

# Slide 2: Introducción
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Introducción a la Aplicación"
tf = slide.placeholders[1].text_frame
tf.text = "GymSystem Móvil es una aplicación Android nativa diseñada para administrar un gimnasio."
p = tf.add_paragraph()
p.text = "• Lenguaje: Java (Android Studio)"
p = tf.add_paragraph()
p.text = "• Base de Datos: MySQL remota/local a través de JDBC"
p = tf.add_paragraph()
p.text = "• Interfaz: Material Design con soporte responsivo"
apply_theme(slide)

# Slide 3: Estructura de Clases (Navegación)
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Navegación y Bienvenida"
tf = slide.placeholders[1].text_frame
tf.text = "SplashActivity:"
p = tf.add_paragraph()
p.text = "Muestra la pantalla de bienvenida con el logo del gimnasio y un diseño atractivo."
p.level = 1
p = tf.add_paragraph()
p.text = "MenuActivity:"
p.level = 0
p = tf.add_paragraph()
p.text = "Panel principal. Contiene tarjetas dinámicas (CardViews) para acceder a los módulos."
p.level = 1
apply_theme(slide)

# Slide 4: Estructura de Clases (Funciones Core)
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Clases de Funcionalidad"
tf = slide.placeholders[1].text_frame
tf.text = "RegistroSocioActivity:"
p = tf.add_paragraph()
p.text = "Captura datos del socio, membresía y permite adjuntar una fotografía."
p.level = 1
p = tf.add_paragraph()
p.text = "ListaSociosActivity:"
p.level = 0
p = tf.add_paragraph()
p.text = "Muestra usuarios registrados, indicando el estado de membresía con colores."
p.level = 1
p = tf.add_paragraph()
p.text = "MarcajeActivity:"
p.level = 0
p = tf.add_paragraph()
p.text = "Registra la entrada y verifica si el socio tiene acceso permitido."
p.level = 1
apply_theme(slide)

# Slide 5: Gestión de Imágenes
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Gestión de Imágenes"
tf = slide.placeholders[1].text_frame
tf.text = "Logo de la Aplicación:"
p = tf.add_paragraph()
p.text = "Se implementó un logotipo principal en el Splash y Menú."
p.level = 1
p = tf.add_paragraph()
p.text = "Foto de Perfil del Socio:"
p.level = 0
p = tf.add_paragraph()
p.text = "Se invoca la galería nativa usando 'ActivityResultLauncher'."
p.level = 1
p = tf.add_paragraph()
p.text = "Al seleccionar, se guarda la ruta (URI) local y se asocia al perfil en BD."
p.level = 1
apply_theme(slide)

# Slide 6: Arquitectura de Base de Datos
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conexión a MySQL"
tf = slide.placeholders[1].text_frame
tf.text = "Clase ConexionMySQL:"
p = tf.add_paragraph()
p.text = "Administra la conexión segura a la BD usando JDBC."
p.level = 1
p = tf.add_paragraph()
p.text = "Hilos Secundarios (Threads):"
p.level = 0
p = tf.add_paragraph()
p.text = "Consultas fuera del hilo principal (UI Thread) para evitar congelamientos."
p.level = 1
p = tf.add_paragraph()
p.text = "Tablas Principales: usuarios, membresias, asistencias."
p.level = 0
apply_theme(slide)

# Slide 7: Conclusión
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conclusión"
tf = slide.placeholders[1].text_frame
tf.text = "GymSystem Móvil cumple con los requerimientos:"
p = tf.add_paragraph()
p.text = "• Refactorización y código limpio."
p.level = 1
p = tf.add_paragraph()
p.text = "• UX mejorada con botones de regreso."
p.level = 1
p = tf.add_paragraph()
p.text = "• Personalización con fotografías."
p.level = 1
p = tf.add_paragraph()
p.text = "• Sincronización robusta con MySQL."
p.level = 1
apply_theme(slide)

prs.save(r'C:\Proyectos Android\GymSystemMovil\Presentacion_GymSystem_Color.pptx')
